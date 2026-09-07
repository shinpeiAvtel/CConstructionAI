from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.schemas.analysis_schema import DocumentAnalysisResult

SHEET_NAMES = [
    "SUMMARY",
    "SOURCE_DOCUMENT",
    "QUOTE_ITEMS",
    "WORK_TYPE_ANALYSIS",
    "PRICE_STATISTICS",
    "LABOR_ANALYSIS",
    "OUTLIERS",
    "RECOMMENDATIONS",
    "VERIFICATION",
    "AUDIT_INFO",
]

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FONT = Font(bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)


def _safe_value(value, fallback="NOT_AVAILABLE"):
    if value is None:
        return fallback
    text = str(value).strip()
    return text if text else fallback


def _safe_float(value):
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _style_header(ws):
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _set_widths(ws, widths):
    for col_index, width in widths.items():
        ws.column_dimensions[get_column_letter(col_index)].width = width


def _apply_numeric_format(ws, column_index, fmt, start_row=2):
    for row in ws.iter_rows(min_row=start_row, min_col=column_index, max_col=column_index):
        for cell in row:
            if cell.value is None or isinstance(cell.value, str):
                continue
            cell.number_format = fmt


def _recommended_total(analysis: DocumentAnalysisResult):
    recommended_total = None
    coverage_text = "Recommended Coverage: 0 / 0 items"
    if not analysis.items:
        return recommended_total, coverage_text

    recommendation_map = {
        str(item.work_type_code or "").strip().upper(): item
        for item in (analysis.price_recommendations or [])
    }

    total = 0.0
    covered_count = 0
    candidate_count = 0
    for item in analysis.items:
        if item.quantity in (None, 0):
            continue
        candidate_count += 1
        if item.work_type_code is None:
            continue
        rec = recommendation_map.get(str(item.work_type_code).strip().upper())
        if rec is None or rec.recommended_unit_price is None:
            continue
        if item.unit and rec.unit and str(item.unit).upper() != str(rec.unit).upper():
            continue
        total += float(item.quantity) * float(rec.recommended_unit_price)
        covered_count += 1

    if candidate_count > 0:
        recommended_total = total
        coverage_text = f"Recommended Coverage: {covered_count} / {candidate_count} items"
    return recommended_total, coverage_text


def _summary_values(analysis: DocumentAnalysisResult):
    doc = analysis.document
    recommended_total, coverage_text = _recommended_total(analysis)
    original_total = _safe_float(getattr(doc, "total", None))
    diff_amount = None
    diff_pct = None
    if original_total is not None and recommended_total is not None:
        diff_amount = recommended_total - original_total
        if original_total not in (None, 0):
            diff_pct = diff_amount / original_total

    confidence = "NOT_AVAILABLE"
    if analysis.price_recommendations:
        confidence = getattr(analysis.price_recommendations[0], "confidence", "NOT_AVAILABLE")

    return {
        "document_id": getattr(doc, "document_id", None),
        "source_filename": getattr(doc, "source_filename", "NOT_AVAILABLE"),
        "source_type": getattr(doc, "source_type", "NOT_AVAILABLE"),
        "vendor": getattr(doc, "vendor_name", "NOT_AVAILABLE"),
        "project": getattr(doc, "project_code", "NOT_AVAILABLE"),
        "quote_number": getattr(doc, "quote_number", "NOT_AVAILABLE"),
        "quote_date": getattr(doc, "quote_date", "NOT_AVAILABLE"),
        "verification_status": getattr(doc, "verification_status", "NOT_AVAILABLE"),
        "analysis_date": getattr(analysis.metadata, "analyzed_at", datetime.now()) or datetime.now(),
        "total_items": getattr(analysis.verification_summary, "total_items", 0),
        "verified_items": getattr(analysis.verification_summary, "verified_items", 0),
        "needs_review_items": getattr(analysis.verification_summary, "needs_review_items", 0),
        "unmapped_items": getattr(analysis.verification_summary, "unmapped_items", 0),
        "rejected_items": getattr(analysis.verification_summary, "rejected_items", 0),
        "analysis_eligible_items": getattr(analysis.verification_summary, "analysis_eligible_items", 0),
        "excluded_items": getattr(analysis.verification_summary, "excluded_items", 0),
        "average_mapping_confidence": getattr(analysis.verification_summary, "average_mapping_confidence", None),
        "original_quote_total": original_total,
        "recommended_estimated_total": recommended_total,
        "difference_amount": diff_amount,
        "difference_pct": diff_pct,
        "recommendation_confidence": confidence,
        "recommended_coverage": coverage_text,
    }


def build_summary_sheet(ws, analysis: DocumentAnalysisResult):
    ws.title = "SUMMARY"
    ws["A1"] = "CConstructionAI"
    ws["A2"] = "Estimate Analysis Report"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"].font = Font(size=12, bold=True)

    summary = _summary_values(analysis)
    left_rows = [
        ("Document ID", summary["document_id"]),
        ("Filename", summary["source_filename"]),
        ("Source Type", summary["source_type"]),
        ("Vendor", summary["vendor"]),
        ("Project", summary["project"]),
        ("Quote Number", summary["quote_number"]),
        ("Quote Date", summary["quote_date"]),
        ("Verification Status", summary["verification_status"]),
        ("Analysis Date", summary["analysis_date"].strftime("%Y-%m-%d %H:%M:%S") if hasattr(summary["analysis_date"], "strftime") else summary["analysis_date"]),
    ]
    for row_index, (label, value) in enumerate(left_rows, start=5):
        ws[f"A{row_index}"] = label
        ws[f"B{row_index}"] = value

    right_rows = [
        ("Total Items", summary["total_items"]),
        ("Verified Items", summary["verified_items"]),
        ("Needs Review Items", summary["needs_review_items"]),
        ("Unmapped Items", summary["unmapped_items"]),
        ("Rejected Items", summary["rejected_items"]),
        ("Analysis Eligible Items", summary["analysis_eligible_items"]),
        ("Excluded Items", summary["excluded_items"]),
        ("Average Mapping Confidence", summary["average_mapping_confidence"]),
    ]
    for row_index, (label, value) in enumerate(right_rows, start=5):
        ws[f"D{row_index}"] = label
        ws[f"E{row_index}"] = value

    pricing_rows = [
        ("Original Quote Total", summary["original_quote_total"]),
        ("Recommended Estimated Total", summary["recommended_estimated_total"]),
        ("Difference Amount", summary["difference_amount"]),
        ("Difference %", summary["difference_pct"]),
        ("Recommendation Confidence", summary["recommendation_confidence"]),
        ("Recommended Coverage", summary["recommended_coverage"]),
    ]
    for row_index, (label, value) in enumerate(pricing_rows, start=16):
        ws[f"A{row_index}"] = label
        ws[f"B{row_index}"] = value

    _set_widths(ws, {1: 22, 2: 24, 4: 24, 5: 18})
    ws.freeze_panes = "A5"
    for cell in ws["1"]:
        cell.font = Font(bold=True, size=12)
    for column in ["A", "B", "D", "E"]:
        for cell in ws[column]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _apply_numeric_format(ws, 2, "#,##0.00", start_row=5)
    _apply_numeric_format(ws, 5, "#,##0.00", start_row=5)
    _apply_numeric_format(ws, 2, "#,##0.00", start_row=16)
    _apply_numeric_format(ws, 2, "0.0%", start_row=18)


def build_source_document_sheet(ws, analysis: DocumentAnalysisResult):
    ws.title = "SOURCE_DOCUMENT"
    doc = analysis.document
    headers = [
        "Document ID",
        "Source Filename",
        "Source Type",
        "Vendor",
        "Project",
        "Quote Number",
        "Quote Date",
        "Currency",
        "Subtotal",
        "Tax",
        "Total",
        "Uploaded At",
        "Verification Status",
    ]
    row = [
        getattr(doc, "document_id", None),
        getattr(doc, "source_filename", "NOT_AVAILABLE"),
        getattr(doc, "source_type", "NOT_AVAILABLE"),
        getattr(doc, "vendor_name", "NOT_AVAILABLE"),
        getattr(doc, "project_code", "NOT_AVAILABLE"),
        getattr(doc, "quote_number", "NOT_AVAILABLE"),
        getattr(doc, "quote_date", "NOT_AVAILABLE"),
        getattr(doc, "currency", "NOT_AVAILABLE"),
        getattr(doc, "subtotal", None),
        getattr(doc, "tax", None),
        getattr(doc, "total", None),
        getattr(doc, "uploaded_at", None),
        getattr(doc, "verification_status", "NOT_AVAILABLE"),
    ]
    ws.append(headers)
    ws.append(row)
    _style_header(ws)
    _set_widths(ws, {1: 16, 2: 24, 3: 18, 4: 18, 5: 18, 6: 18, 7: 18, 8: 12, 9: 14, 10: 12, 11: 12, 12: 20, 13: 18})
    _apply_numeric_format(ws, 9, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 10, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 11, "#,##0.00", start_row=2)


def build_quote_items_sheet(ws, analysis: DocumentAnalysisResult):
    ws.title = "QUOTE_ITEMS"
    headers = [
        "Line",
        "Raw Description",
        "Normalized Description",
        "Work Type Code",
        "Work Type Name",
        "Category",
        "Sub Category",
        "Manufacturer",
        "Model",
        "Quantity",
        "Unit",
        "Source Unit Price",
        "Calculated Unit Price",
        "Normalized Unit Price",
        "Source Amount",
        "Calculated Amount",
        "Quoted Person Days",
        "Quoted Labor Amount",
        "Mapping Confidence",
        "Verification Status",
        "Outlier",
        "Outlier Reason",
    ]
    ws.append(headers)
    for item in analysis.items or []:
        ws.append([
            getattr(item, "line_number", None),
            getattr(item, "raw_description", "NOT_AVAILABLE"),
            getattr(item, "normalized_description", "NOT_AVAILABLE"),
            getattr(item, "work_type_code", "NOT_AVAILABLE"),
            getattr(item, "work_type_name", "NOT_AVAILABLE"),
            getattr(item, "category", "NOT_AVAILABLE"),
            getattr(item, "sub_category", "NOT_AVAILABLE"),
            getattr(item, "manufacturer", "NOT_AVAILABLE"),
            getattr(item, "model_number", "NOT_AVAILABLE"),
            getattr(item, "quantity", None),
            getattr(item, "unit", "NOT_AVAILABLE"),
            getattr(item, "source_unit_price", None),
            getattr(item, "calculated_unit_price", None),
            getattr(item, "normalized_unit_price", None),
            getattr(item, "source_amount", None),
            getattr(item, "calculated_amount", None),
            getattr(item, "quoted_person_days", None),
            getattr(item, "quoted_labor_amount", None),
            getattr(item, "mapping_confidence", None),
            getattr(item, "verification_status", "NOT_AVAILABLE"),
            "TRUE" if getattr(item, "is_outlier", False) else "FALSE",
            getattr(item, "outlier_reason", "NOT_AVAILABLE"),
        ])
    _style_header(ws)
    _apply_numeric_format(ws, 10, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 12, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 13, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 15, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 16, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 17, "0.0000", start_row=2)
    _apply_numeric_format(ws, 18, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 19, "0.0%", start_row=2)
    _set_widths(ws, {1: 10, 2: 30, 3: 26, 4: 16, 5: 18, 6: 18, 7: 18, 8: 18, 9: 16, 10: 12, 11: 12, 12: 16, 13: 16, 14: 18, 15: 16, 16: 16, 17: 16, 18: 18, 19: 14, 20: 18, 21: 12, 22: 22})


def build_work_type_analysis_sheet(ws, analysis: DocumentAnalysisResult):
    ws.title = "WORK_TYPE_ANALYSIS"
    headers = [
        "Work Type Code",
        "Work Type Name",
        "Unit",
        "Sample Count",
        "Verified Samples",
        "Unverified Samples",
        "Excluded Samples",
        "Mean",
        "Median",
        "Weighted Mean",
        "Std Dev",
        "Min",
        "Max",
        "Q1",
        "Q3",
        "IQR",
        "Lower Bound",
        "Upper Bound",
        "Outlier Count",
        "Outlier Ratio",
        "Recent Median",
    ]
    ws.append(headers)
    for item in analysis.work_type_analysis or []:
        ws.append([
            getattr(item, "work_type_code", "NOT_AVAILABLE"),
            getattr(item, "work_type_name", "NOT_AVAILABLE"),
            getattr(item, "unit", "NOT_AVAILABLE"),
            getattr(item, "sample_count", 0),
            getattr(item, "verified_sample_count", 0),
            getattr(item, "unverified_sample_count", 0),
            getattr(item, "excluded_sample_count", 0),
            getattr(item, "mean", None),
            getattr(item, "median", None),
            getattr(item, "weighted_mean", None),
            getattr(item, "standard_deviation", None),
            getattr(item, "minimum", None),
            getattr(item, "maximum", None),
            getattr(item, "q1", None),
            getattr(item, "q3", None),
            getattr(item, "iqr", None),
            getattr(item, "lower_bound", None),
            getattr(item, "upper_bound", None),
            getattr(item, "outlier_count", 0),
            getattr(item, "outlier_ratio", None),
            getattr(item, "recent_median", None),
        ])
    _style_header(ws)
    _apply_numeric_format(ws, 4, "#,##0", start_row=2)
    _apply_numeric_format(ws, 5, "#,##0", start_row=2)
    _apply_numeric_format(ws, 6, "#,##0", start_row=2)
    _apply_numeric_format(ws, 7, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 8, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 9, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 10, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 11, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 12, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 13, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 14, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 15, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 16, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 17, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 18, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 19, "#,##0", start_row=2)
    _apply_numeric_format(ws, 20, "0.0%", start_row=2)
    _apply_numeric_format(ws, 21, "#,##0.00", start_row=2)


def build_price_statistics_sheet(ws, analysis: DocumentAnalysisResult):
    ws.title = "PRICE_STATISTICS"
    headers = ["Work Type", "Unit", "Sample Count", "Mean", "Median", "Weighted Mean", "Q1", "Q3", "Recent Median"]
    ws.append(headers)
    for item in analysis.work_type_analysis or []:
        ws.append([
            getattr(item, "work_type_code", "NOT_AVAILABLE"),
            getattr(item, "unit", "NOT_AVAILABLE"),
            getattr(item, "sample_count", 0),
            getattr(item, "mean", None),
            getattr(item, "median", None),
            getattr(item, "weighted_mean", None),
            getattr(item, "q1", None),
            getattr(item, "q3", None),
            getattr(item, "recent_median", None),
        ])
    _style_header(ws)
    _apply_numeric_format(ws, 3, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 4, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 5, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 6, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 7, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 8, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 9, "#,##0.00", start_row=2)
    _set_widths(ws, {1: 18, 2: 12, 3: 14, 4: 12, 5: 12, 6: 14, 7: 12, 8: 12, 9: 14})


def build_labor_analysis_sheet(ws, analysis: DocumentAnalysisResult):
    ws.title = "LABOR_ANALYSIS"
    headers = [
        "Work Type Code",
        "Labor Code",
        "Unit",
        "Current Standard Labor",
        "Quoted Labor Sample Count",
        "Quoted Labor Mean",
        "Quoted Labor Median",
        "Actual Work Sample Count",
        "Actual Work Mean",
        "Actual Work Median",
        "Recommended Labor Candidate",
        "Difference From Current %",
        "Confidence",
        "Human Review Required",
        "Reason",
    ]
    ws.append(headers)
    for item in analysis.labor_analysis or []:
        ws.append([
            getattr(item, "work_type_code", "NOT_AVAILABLE"),
            getattr(item, "labor_code", "NOT_AVAILABLE"),
            getattr(item, "unit", "NOT_AVAILABLE"),
            getattr(item, "current_standard_labor", None),
            getattr(item, "quoted_labor_sample_count", 0),
            getattr(item, "quoted_labor_mean", None),
            getattr(item, "quoted_labor_median", None),
            getattr(item, "actual_work_sample_count", 0),
            getattr(item, "actual_work_mean", None),
            getattr(item, "actual_work_median", None),
            getattr(item, "recommended_labor_candidate", None),
            getattr(item, "difference_from_current_percent", None),
            getattr(item, "confidence", "NOT_AVAILABLE"),
            "TRUE" if getattr(item, "human_review_required", False) else "FALSE",
            getattr(item, "reason", "NOT_AVAILABLE"),
        ])
    _style_header(ws)
    _apply_numeric_format(ws, 4, "0.0000", start_row=2)
    _apply_numeric_format(ws, 5, "#,##0", start_row=2)
    _apply_numeric_format(ws, 6, "0.0000", start_row=2)
    _apply_numeric_format(ws, 7, "0.0000", start_row=2)
    _apply_numeric_format(ws, 8, "#,##0", start_row=2)
    _apply_numeric_format(ws, 9, "0.0000", start_row=2)
    _apply_numeric_format(ws, 10, "0.0000", start_row=2)
    _apply_numeric_format(ws, 11, "0.0000", start_row=2)
    _apply_numeric_format(ws, 12, "0.0%", start_row=2)
    _set_widths(ws, {1: 18, 2: 18, 3: 12, 4: 20, 5: 18, 6: 18, 7: 18, 8: 20, 9: 16, 10: 18, 11: 18, 12: 22, 13: 14, 14: 18, 15: 48})


def build_outliers_sheet(ws, analysis: DocumentAnalysisResult):
    ws.title = "OUTLIERS"
    headers = ["Line", "Raw Description", "Normalized Description", "Work Type Code", "Unit", "Analysis Unit Price", "Outlier Reason", "Verification Status", "Q1", "Q3", "Lower Bound", "Upper Bound"]
    ws.append(headers)
    for item in analysis.outliers or []:
        ws.append([
            getattr(item, "line_number", None),
            getattr(item, "raw_description", "NOT_AVAILABLE"),
            getattr(item, "normalized_description", "NOT_AVAILABLE"),
            getattr(item, "work_type_code", "NOT_AVAILABLE"),
            getattr(item, "unit", "NOT_AVAILABLE"),
            getattr(item, "normalized_unit_price", None) or getattr(item, "calculated_unit_price", None),
            getattr(item, "outlier_reason", "NOT_AVAILABLE"),
            getattr(item, "verification_status", "NOT_AVAILABLE"),
            None,
            None,
            None,
            None,
        ])
    _style_header(ws)
    _apply_numeric_format(ws, 6, "#,##0.00", start_row=2)
    _set_widths(ws, {1: 10, 2: 30, 3: 28, 4: 16, 5: 12, 6: 18, 7: 24, 8: 18, 9: 12, 10: 12, 11: 14, 12: 14})


def build_recommendations_sheet(ws, analysis: DocumentAnalysisResult):
    ws.title = "RECOMMENDATIONS"
    headers = [
        "Work Type Code",
        "Unit",
        "Historical Mean",
        "Historical Median",
        "Recent Median",
        "Recommended Unit Price",
        "LOW",
        "BASE",
        "HIGH",
        "Sample Count",
        "Vendor Count",
        "Outlier Ratio",
        "Confidence",
        "Confidence Score",
        "Human Review Required",
        "Reason",
    ]
    ws.append(headers)
    for item in analysis.price_recommendations or []:
        ws.append([
            getattr(item, "work_type_code", "NOT_AVAILABLE"),
            getattr(item, "unit", "NOT_AVAILABLE"),
            getattr(item, "historical_mean", None),
            getattr(item, "historical_median", None),
            getattr(item, "recent_median", None),
            getattr(item, "recommended_unit_price", None),
            getattr(item, "low_unit_price", None),
            getattr(item, "base_unit_price", None),
            getattr(item, "high_unit_price", None),
            getattr(item, "sample_count", 0),
            getattr(item, "vendor_count", None),
            getattr(item, "outlier_ratio", None),
            getattr(item, "confidence", "NOT_AVAILABLE"),
            getattr(item, "confidence_score", None),
            "TRUE" if getattr(item, "human_review_required", False) else "FALSE",
            getattr(item, "reason", "NOT_AVAILABLE"),
        ])
    _style_header(ws)
    _apply_numeric_format(ws, 3, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 4, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 5, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 6, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 7, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 8, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 9, "#,##0.00", start_row=2)
    _apply_numeric_format(ws, 10, "#,##0", start_row=2)
    _apply_numeric_format(ws, 11, "#,##0", start_row=2)
    _apply_numeric_format(ws, 12, "0.0%", start_row=2)
    _apply_numeric_format(ws, 14, "0.0%", start_row=2)
    _set_widths(ws, {1: 18, 2: 12, 3: 16, 4: 18, 5: 18, 6: 20, 7: 12, 8: 12, 9: 12, 10: 14, 11: 12, 12: 12, 13: 16, 14: 20, 15: 18, 16: 48})


def build_verification_sheet(ws, analysis: DocumentAnalysisResult):
    ws.title = "VERIFICATION"
    headers = [
        "Line",
        "Raw Description",
        "Normalized Description",
        "Work Type Code",
        "Mapping Confidence",
        "Verification Status",
        "Needs Review",
        "Outlier",
    ]
    ws.append(headers)
    for item in analysis.items or []:
        ws.append([
            getattr(item, "line_number", None),
            getattr(item, "raw_description", "NOT_AVAILABLE"),
            getattr(item, "normalized_description", "NOT_AVAILABLE"),
            getattr(item, "work_type_code", "NOT_AVAILABLE"),
            getattr(item, "mapping_confidence", None),
            getattr(item, "verification_status", "NOT_AVAILABLE"),
            "TRUE" if getattr(item, "verification_status", "").upper() == "NEEDS_REVIEW" else "FALSE",
            "TRUE" if getattr(item, "is_outlier", False) else "FALSE",
        ])

    summary = analysis.verification_summary
    summary_start = ws.max_row + 3
    ws[f"A{summary_start}"] = "Summary"
    ws[f"A{summary_start}"] .font = Font(bold=True)
    summary_rows = [
        ("Total", getattr(summary, "total_items", 0)),
        ("Verified", getattr(summary, "verified_items", 0)),
        ("Needs Review", getattr(summary, "needs_review_items", 0)),
        ("Unmapped", getattr(summary, "unmapped_items", 0)),
        ("Rejected", getattr(summary, "rejected_items", 0)),
        ("Eligible", getattr(summary, "analysis_eligible_items", 0)),
        ("Excluded", getattr(summary, "excluded_items", 0)),
    ]
    for idx, (label, value) in enumerate(summary_rows, start=summary_start + 1):
        ws[f"A{idx}"] = label
        ws[f"B{idx}"] = value

    _style_header(ws)
    _apply_numeric_format(ws, 5, "0.0%", start_row=2)
    _set_widths(ws, {1: 10, 2: 30, 3: 28, 4: 18, 5: 14, 6: 18, 7: 14, 8: 12})


def build_audit_info_sheet(ws, analysis: DocumentAnalysisResult):
    ws.title = "AUDIT_INFO"
    metadata = analysis.metadata
    row = [
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        getattr(analysis.document, "document_id", None),
        getattr(metadata, "analysis_version", "NOT_AVAILABLE"),
        getattr(metadata, "analyzed_at", datetime.now()) or datetime.now(),
        getattr(metadata, "verification_filter", "NOT_AVAILABLE"),
        getattr(metadata, "source_type_filter", "NOT_AVAILABLE"),
        getattr(metadata, "vendor_filter", "NOT_AVAILABLE"),
        getattr(metadata, "project_filter", "NOT_AVAILABLE"),
        getattr(metadata, "included_samples", 0),
        getattr(metadata, "excluded_samples", 0),
        getattr(metadata, "outlier_method", "NOT_AVAILABLE"),
        "NOT_AVAILABLE",
    ]
    ws.append([
        "Exported At",
        "Document ID",
        "Analysis Version",
        "Analysis Date",
        "Verification Filter",
        "Source Type Filter",
        "Vendor Filter",
        "Project Filter",
        "Included Samples",
        "Excluded Samples",
        "Outlier Method",
        "Exported By",
    ])
    ws.append(row)
    _style_header(ws)
    _apply_numeric_format(ws, 2, "#,##0", start_row=2)
    _apply_numeric_format(ws, 9, "#,##0", start_row=2)
    _apply_numeric_format(ws, 10, "#,##0", start_row=2)
    _set_widths(ws, {1: 20, 2: 14, 3: 18, 4: 20, 5: 18, 6: 18, 7: 18, 8: 18, 9: 18, 10: 18, 11: 18, 12: 18})


def build_workbook(analysis: DocumentAnalysisResult) -> Workbook:
    wb = Workbook()
    wb.active.title = "SUMMARY"
    for name in SHEET_NAMES[1:]:
        wb.create_sheet(title=name)

    sheets = {name: wb[name] for name in SHEET_NAMES}
    build_summary_sheet(sheets["SUMMARY"], analysis)
    build_source_document_sheet(sheets["SOURCE_DOCUMENT"], analysis)
    build_quote_items_sheet(sheets["QUOTE_ITEMS"], analysis)
    build_work_type_analysis_sheet(sheets["WORK_TYPE_ANALYSIS"], analysis)
    build_price_statistics_sheet(sheets["PRICE_STATISTICS"], analysis)
    build_labor_analysis_sheet(sheets["LABOR_ANALYSIS"], analysis)
    build_outliers_sheet(sheets["OUTLIERS"], analysis)
    build_recommendations_sheet(sheets["RECOMMENDATIONS"], analysis)
    build_verification_sheet(sheets["VERIFICATION"], analysis)
    build_audit_info_sheet(sheets["AUDIT_INFO"], analysis)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = True
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    cell.value = ""
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = THIN_BORDER
        ws.auto_filter.ref = ws.dimensions
    return wb


def build_filename(document_id: int | str, analysis_date: str | datetime | None = None) -> str:
    if analysis_date is None:
        analysis_date = datetime.now()
    if isinstance(analysis_date, datetime):
        date_stamp = analysis_date.strftime("%Y%m%d")
    else:
        date_stamp = str(analysis_date).replace("-", "").replace("/", "")[:8]
    sanitized_doc = re.sub(r"[^A-Za-z0-9_]", "", str(document_id))
    sanitized_doc = sanitized_doc or "DOCUMENT"
    return f"ESTIMATE_ANALYSIS_{sanitized_doc}_{date_stamp}.xlsx"


def build_workbook_bytes(analysis: DocumentAnalysisResult) -> bytes:
    workbook = build_workbook(analysis)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


__all__ = [
    "SHEET_NAMES",
    "build_workbook",
    "build_summary_sheet",
    "build_source_document_sheet",
    "build_quote_items_sheet",
    "build_work_type_analysis_sheet",
    "build_price_statistics_sheet",
    "build_labor_analysis_sheet",
    "build_outliers_sheet",
    "build_recommendations_sheet",
    "build_verification_sheet",
    "build_audit_info_sheet",
    "build_filename",
    "build_workbook_bytes",
]
