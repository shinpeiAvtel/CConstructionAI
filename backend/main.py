from contextlib import asynccontextmanager
from datetime import datetime
from statistics import mean, median, quantiles

from fastapi import Depends, FastAPI, HTTPException, Request, UploadFile, File, Form
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlmodel import Session, select

from backend.database import create_db_and_tables, get_session
from backend.models import (
    ActualWorkRecord,
    ActualWorkUpdate,
    AuditLog,
    CombinedEstimateRequest,
    CompanySetting,
    Estimate,
    EstimateCreateRequest,
    EstimateItem,
    EstimateItemCreateRequest,
    EstimateItemUpdateRequest,
    EstimateStatusUpdateRequest,
    LaborCostMaster,
    LaborEstimateRequest,
    LaborMaster,
    MaterialMaster,
    Project,
    StandardRevisionCandidate,
    StandardRevisionReviewRequest,
    UploadedEstimate,
)

from io import BytesIO
from pathlib import Path

from fastapi.responses import StreamingResponse, HTMLResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from backend.routers.work_type_router import router as work_type_router
from backend.routers.estimate_document_router import router as estimate_document_router

# =========================================================
# APPLICATION STARTUP
# =========================================================

@asynccontextmanager
async def lifespan(app: FastAPI):
    create_db_and_tables()
    yield


# =========================================================
# DATA QUALITY SETTINGS
# =========================================================

MIN_ANALYSIS_SAMPLE_COUNT = 5

RECOMMENDED_ANALYSIS_SAMPLE_COUNT = 10

OUTLIER_IQR_MULTIPLIER = 1.5

# =========================================================
# STANDARD REVISION SETTINGS
# =========================================================

REVISION_MIN_SAMPLE_COUNT = 10

REVISION_DIFFERENCE_THRESHOLD_PERCENT = 15.0

REVISION_MAX_MEAN_MEDIAN_GAP_PERCENT = 20.0

REVISION_MAX_OUTLIER_RATIO = 0.25


# =========================================================
# DATE VALIDATION
# =========================================================

def validate_date_string(
    value: str,
    field_name: str,
):
    try:
        datetime.strptime(
            value,
            "%Y-%m-%d",
        )
    except ValueError:
        raise HTTPException(
            status_code=400,
            detail=(
                f"{field_name} must be "
                "YYYY-MM-DD format."
            ),
        )

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
TEMPLATE_DIR = BASE_DIR / "templates"

app = FastAPI(
    title="Construction Estimator",
    description="工事歩掛・積算システム",
    version="0.3.0",
    lifespan=lifespan,
)

app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
templates = Jinja2Templates(directory=str(TEMPLATE_DIR))

app.include_router(work_type_router)
app.include_router(estimate_document_router)


# =========================================================
# SYSTEM
# =========================================================

@app.get("/")
def root():
    return {
        "system": "Construction Estimator",
        "status": "running",
        "version": "0.3.0",
    }


@app.get("/health")
def health_check():
    return {
        "status": "ok"
    }


@app.get("/mobile", response_class=HTMLResponse)
async def mobile_estimate_analysis(request: Request):
    return templates.TemplateResponse(
        request,
        "mobile_analysis.html",
        {},
    )


@app.get("/upload", response_class=HTMLResponse)
async def upload_mobile_ui(request: Request):
    return templates.TemplateResponse(
        request,
        "upload_mobile.html",
        {},
    )


# =========================================================
# UPLOADS
# =========================================================


@app.post("/upload-estimate")
async def upload_estimate(
    file: UploadFile = File(...),
    uploader: str | None = Form(None),
    project_code: str | None = Form(None),
    session: Session = Depends(get_session),
):
    if file.content_type != "application/pdf":
        raise HTTPException(
            status_code=400,
            detail="Only PDF files are accepted.",
        )

    uploads_dir = Path(__file__).resolve().parent.parent / "data" / "uploads"
    uploads_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.utcnow().strftime("%Y%m%d%H%M%S%f")
    safe_name = f"{timestamp}_{file.filename}"
    dest = uploads_dir / safe_name

    contents = await file.read()
    dest.write_bytes(contents)

    record = UploadedEstimate(
        filename=safe_name,
        original_filename=file.filename,
        uploader=uploader,
        project_code=project_code,
        file_path=str(dest.as_posix()),
    )

    session.add(record)
    session.commit()
    session.refresh(record)

    return {
        "id": record.id,
        "filename": record.filename,
        "original_filename": record.original_filename,
        "uploader": record.uploader,
        "project_code": record.project_code,
        "uploaded_at": record.uploaded_at,
    }


@app.get("/upload-page")
def upload_page():
    html = """
    <html>
      <head>
        <title>見積りPDFアップロード</title>
      </head>
      <body>
        <h1>見積りPDFアップロード</h1>
        <form action="/upload-estimate" enctype="multipart/form-data" method="post">
          <label>営業担当者名: <input type="text" name="uploader" /></label><br/>
          <label>プロジェクトコード: <input type="text" name="project_code" /></label><br/>
          <input name="file" type="file" accept="application/pdf" /><br/>
          <button type="submit">アップロード</button>
        </form>
      </body>
    </html>
    """

    return HTMLResponse(content=html, status_code=200)


@app.get("/uploads")
def list_uploads(limit: int = 50, session: Session = Depends(get_session)):
    items = session.exec(select(UploadedEstimate).order_by(UploadedEstimate.uploaded_at.desc())).all()
    return items[:limit]


# =========================================================
# LABOR MASTER
# =========================================================

@app.post(
    "/labor-master",
    response_model=LaborMaster,
)
def create_labor_master(
    item: LaborMaster,
    session: Session = Depends(get_session),
):
    existing = session.exec(
        select(LaborMaster).where(
            LaborMaster.code == item.code
        )
    ).first()

    if existing:
        raise HTTPException(
            status_code=409,
            detail="This labor code already exists.",
        )

    session.add(item)
    session.commit()
    session.refresh(item)

    return item


@app.get("/labor-master")
def read_labor_master(
    session: Session = Depends(get_session),
):
    items = session.exec(
        select(LaborMaster).order_by(
            LaborMaster.code
        )
    ).all()

    return items


# =========================================================
# COMPANY SETTINGS
# =========================================================

@app.post(
    "/company-settings",
    response_model=CompanySetting,
)
def create_company_setting(
    item: CompanySetting,
    session: Session = Depends(get_session),
):
    existing = session.exec(
        select(CompanySetting).where(
            CompanySetting.setting_code
            == item.setting_code
        )
    ).first()

    if existing:
        raise HTTPException(
            status_code=409,
            detail="This setting code already exists.",
        )

    session.add(item)
    session.commit()
    session.refresh(item)

    return item


@app.get("/company-settings")
def read_company_settings(
    session: Session = Depends(get_session),
):
    items = session.exec(
        select(CompanySetting).order_by(
            CompanySetting.setting_code
        )
    ).all()

    return items


# =========================================================
# LABOR COST MASTER
# =========================================================

@app.post(
    "/labor-cost-master",
    response_model=LaborCostMaster,
)
def create_labor_cost_master(
    item: LaborCostMaster,
    session: Session = Depends(get_session),
):
    existing = session.exec(
        select(LaborCostMaster).where(
            LaborCostMaster.cost_code
            == item.cost_code
        )
    ).first()

    if existing:
        raise HTTPException(
            status_code=409,
            detail="This cost code already exists.",
        )

    session.add(item)
    session.commit()
    session.refresh(item)

    return item


@app.get("/labor-cost-master")
def read_labor_cost_master(
    session: Session = Depends(get_session),
):
    items = session.exec(
        select(LaborCostMaster).order_by(
            LaborCostMaster.cost_code
        )
    ).all()

    return items


# =========================================================
# PROJECT
# =========================================================

@app.post(
    "/projects",
    response_model=Project,
)
def create_project(
    item: Project,
    session: Session = Depends(get_session),
):
    existing = session.exec(
        select(Project).where(
            Project.project_code
            == item.project_code
        )
    ).first()

    if existing:
        raise HTTPException(
            status_code=409,
            detail="Project code already exists.",
        )

    session.add(item)
    session.commit()
    session.refresh(item)

    return item


@app.get("/projects")
def read_projects(
    session: Session = Depends(get_session),
):
    items = session.exec(
        select(Project).order_by(
            Project.project_code
        )
    ).all()

    return items

# =========================================================
# ACTUAL WORK RECORD
# =========================================================

@app.post("/actual-work")
def create_actual_work(
    item: ActualWorkRecord,
    session: Session = Depends(get_session),
):

    project = session.exec(
        select(Project).where(
            Project.project_code
            == item.project_code
        )
    ).first()

    if not project:
        raise HTTPException(
            status_code=404,
            detail="Project code not found.",
        )

    labor_item = session.exec(
        select(LaborMaster).where(
            LaborMaster.code
            == item.labor_code
        )
    ).first()

    if not labor_item:
        raise HTTPException(
            status_code=404,
            detail="Labor master code not found.",
        )

    if item.quantity <= 0:
        raise HTTPException(
            status_code=400,
            detail="Quantity must be greater than zero.",
        )

    if item.workers <= 0:
        raise HTTPException(
            status_code=400,
            detail="Workers must be greater than zero.",
        )

    if item.work_minutes <= 0:
        raise HTTPException(
            status_code=400,
            detail="Work minutes must be greater than zero.",
        )

    company_setting = session.exec(
        select(CompanySetting).where(
            CompanySetting.setting_code
            == "MINUTES_PER_PERSON_DAY"
        )
    ).first()

    if not company_setting:
        raise HTTPException(
            status_code=404,
            detail=(
                "MINUTES_PER_PERSON_DAY "
                "company setting not found."
            ),
        )

    actual_person_days = (
        item.workers
        * item.work_minutes
        / company_setting.value
    )

    actual_labor_per_unit = (
        actual_person_days
        / item.quantity
    )


        # -----------------------------------------------------
    # DUPLICATE CHECK
    # -----------------------------------------------------

    duplicate = session.exec(
        select(ActualWorkRecord).where(
            ActualWorkRecord.project_code
            == item.project_code,

            ActualWorkRecord.labor_code
            == item.labor_code,

            ActualWorkRecord.work_date
            == item.work_date,

            ActualWorkRecord.quantity
            == item.quantity,

            ActualWorkRecord.workers
            == item.workers,

            ActualWorkRecord.work_minutes
            == item.work_minutes,

            ActualWorkRecord.site_condition
            == item.site_condition,
        )
    ).first()

    if duplicate:
        raise HTTPException(
            status_code=409,
            detail=(
                "Possible duplicate actual work "
                f"record. Existing ID: {duplicate.id}"
            ),
        )

    session.add(item)
    session.commit()
    session.refresh(item)

    return {
        "id":
            item.id,

        "project_code":
            item.project_code,

        "labor_code":
            item.labor_code,

        "work_date":
            item.work_date,

        "quantity":
            item.quantity,

        "workers":
            item.workers,

        "work_minutes":
            item.work_minutes,

        "actual_person_days":
            round(
                actual_person_days,
                4,
            ),

        "actual_labor_per_unit":
            round(
                actual_labor_per_unit,
                4,
            ),

        "site_condition":
            item.site_condition,

        "description":
            item.description,
    }


@app.get("/actual-work")
def read_actual_work(
    project_code: str | None = None,
    labor_code: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    site_condition: str | None = None,
    session: Session = Depends(get_session),
):

    if start_date:
        validate_date_string(
            start_date,
            "start_date",
        )

    if end_date:
        validate_date_string(
            end_date,
            "end_date",
        )

    if (
        start_date
        and end_date
        and start_date > end_date
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "start_date must not be "
                "later than end_date."
            ),
        )

    statement = select(
        ActualWorkRecord
    )

    if project_code:
        statement = statement.where(
            ActualWorkRecord.project_code
            == project_code
        )

    if labor_code:
        statement = statement.where(
            ActualWorkRecord.labor_code
            == labor_code
        )

    if start_date:
        statement = statement.where(
            ActualWorkRecord.work_date
            >= start_date
        )

    if end_date:
        statement = statement.where(
            ActualWorkRecord.work_date
            <= end_date
        )

    if site_condition:
        statement = statement.where(
            ActualWorkRecord.site_condition
            == site_condition
        )

    statement = statement.order_by(
        ActualWorkRecord.work_date,
        ActualWorkRecord.id,
    )

    items = session.exec(
        statement
    ).all()

    return {
        "count": len(items),
        "filters": {
            "project_code": project_code,
            "labor_code": labor_code,
            "start_date": start_date,
            "end_date": end_date,
            "site_condition": site_condition,
        },
        "records": items,
    }




    # -----------------------------------------------------
    # 2. LABOR MASTER CHECK
    # -----------------------------------------------------

    labor_item = session.exec(
        select(LaborMaster).where(
            LaborMaster.code == item.labor_code
        )
    ).first()

    if not labor_item:
        raise HTTPException(
            status_code=404,
            detail="Labor master code not found.",
        )

    # -----------------------------------------------------
    # 3. INPUT VALIDATION
    # -----------------------------------------------------

    validate_date_string(
        item.work_date,
        "work_date",
    )


    if item.quantity <= 0:
        raise HTTPException(
            status_code=400,
            detail="Quantity must be greater than zero.",
        )

    if item.workers <= 0:
        raise HTTPException(
            status_code=400,
            detail="Workers must be greater than zero.",
        )

    if item.work_minutes <= 0:
        raise HTTPException(
            status_code=400,
            detail="Work minutes must be greater than zero.",
        )

    # -----------------------------------------------------
    # 4. COMPANY SETTING
    # -----------------------------------------------------

    company_setting = session.exec(
        select(CompanySetting).where(
            CompanySetting.setting_code
            == "MINUTES_PER_PERSON_DAY"
        )
    ).first()

    if not company_setting:
        raise HTTPException(
            status_code=404,
            detail=(
                "MINUTES_PER_PERSON_DAY "
                "company setting not found."
            ),
        )

    if company_setting.value <= 0:
        raise HTTPException(
            status_code=500,
            detail=(
                "MINUTES_PER_PERSON_DAY "
                "must be greater than zero."
            ),
        )

    # -----------------------------------------------------
    # 5. CALCULATE ACTUAL LABOR
    # -----------------------------------------------------

    actual_person_days = (
        item.workers
        * item.work_minutes
        / company_setting.value
    )

    actual_labor_per_unit = (
        actual_person_days
        / item.quantity
    )

    # -----------------------------------------------------
    # 6. SAVE RAW DATA
    # -----------------------------------------------------

    session.add(item)
    session.commit()
    session.refresh(item)

    # -----------------------------------------------------
    # 7. RESPONSE
    # -----------------------------------------------------

    return {
        "id": item.id,
        "project_code": item.project_code,
        "labor_code": item.labor_code,
        "work_date": item.work_date,
        "quantity": item.quantity,
        "workers": item.workers,
        "work_minutes": item.work_minutes,
        "minutes_per_person_day": company_setting.value,
        "actual_person_days": round(actual_person_days, 4),
        "actual_labor_per_unit": round(
            actual_labor_per_unit,
            4,
        ),
        "site_condition": item.site_condition,
        "description": item.description,
    }




# =========================================================
# ACTUAL WORK UPDATE
# =========================================================

@app.put("/actual-work/{record_id}")
def update_actual_work(
    record_id: int,
    update_data: ActualWorkUpdate,
    session: Session = Depends(get_session),
):

    record = session.get(
        ActualWorkRecord,
        record_id,
    )

    if not record:
        raise HTTPException(
            status_code=404,
            detail="Actual work record not found.",
        )

    update_values = update_data.model_dump(
        exclude_unset=True
    )

    if "work_date" in update_values:
        validate_date_string(
            update_values["work_date"],
            "work_date",
        )

    if (
        "quantity" in update_values
        and update_values["quantity"] <= 0
    ):
        raise HTTPException(
            status_code=400,
            detail="Quantity must be greater than zero.",
        )

    if (
        "workers" in update_values
        and update_values["workers"] <= 0
    ):
        raise HTTPException(
            status_code=400,
            detail="Workers must be greater than zero.",
        )

    if (
        "work_minutes" in update_values
        and update_values["work_minutes"] <= 0
    ):
        raise HTTPException(
            status_code=400,
            detail="Work minutes must be greater than zero.",
        )

    new_project_code = update_values.get(
        "project_code",
        record.project_code,
    )

    new_labor_code = update_values.get(
        "labor_code",
        record.labor_code,
    )

    project = session.exec(
        select(Project).where(
            Project.project_code == new_project_code
        )
    ).first()

    if not project:
        raise HTTPException(
            status_code=404,
            detail="Project code not found.",
        )

    labor_item = session.exec(
        select(LaborMaster).where(
            LaborMaster.code == new_labor_code
        )
    ).first()

    if not labor_item:
        raise HTTPException(
            status_code=404,
            detail="Labor master code not found.",
        )

    for key, value in update_values.items():
        setattr(record, key, value)

    duplicate = session.exec(
        select(ActualWorkRecord).where(
            ActualWorkRecord.id != record.id,
            ActualWorkRecord.project_code == record.project_code,
            ActualWorkRecord.labor_code == record.labor_code,
            ActualWorkRecord.work_date == record.work_date,
            ActualWorkRecord.quantity == record.quantity,
            ActualWorkRecord.workers == record.workers,
            ActualWorkRecord.work_minutes == record.work_minutes,
            ActualWorkRecord.site_condition == record.site_condition,
        )
    ).first()

    if duplicate:
        session.rollback()

        raise HTTPException(
            status_code=409,
            detail="Update would create a duplicate actual work record.",
        )

    session.add(record)
    session.commit()
    session.refresh(record)

    return record
# =========================================================
# ACTUAL WORK DELETE
# =========================================================

@app.delete("/actual-work/{record_id}")
def delete_actual_work(
    record_id: int,
    session: Session = Depends(get_session),
):

    record = session.get(
        ActualWorkRecord,
        record_id,
    )

    if not record:
        raise HTTPException(
            status_code=404,
            detail="Actual work record not found.",
        )

    deleted_record = {
        "id": record.id,
        "project_code": record.project_code,
        "labor_code": record.labor_code,
        "work_date": record.work_date,
    }

    session.delete(record)
    session.commit()

    return {
        "message": (
            "Actual work record deleted."
        ),
        "deleted_record": deleted_record,
    }




def read_actual_work(
    project_code: str | None = None,
    labor_code: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    site_condition: str | None = None,
    session: Session = Depends(get_session),
):

    if start_date:
        validate_date_string(
            start_date,
            "start_date",
        )

    if end_date:
        validate_date_string(
            end_date,
            "end_date",
        )

    if (
        start_date
        and end_date
        and start_date > end_date
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "start_date must not be "
                "later than end_date."
            ),
        )

    statement = select(
        ActualWorkRecord
    )

    if project_code:
        statement = statement.where(
            ActualWorkRecord.project_code
            == project_code
        )

    if labor_code:
        statement = statement.where(
            ActualWorkRecord.labor_code
            == labor_code
        )

    if start_date:
        statement = statement.where(
            ActualWorkRecord.work_date
            >= start_date
        )

    if end_date:
        statement = statement.where(
            ActualWorkRecord.work_date
            <= end_date
        )

    if site_condition:
        statement = statement.where(
            ActualWorkRecord.site_condition
            == site_condition
        )

    statement = statement.order_by(
        ActualWorkRecord.work_date,
        ActualWorkRecord.id,
    )

    items = session.exec(
        statement
    ).all()

    return {
        "count": len(items),
        "filters": {
            "project_code": project_code,
            "labor_code": labor_code,
            "start_date": start_date,
            "end_date": end_date,
            "site_condition": site_condition,
        },
        "records": items,
    }


# =========================================================
# LABOR ANALYSIS
# =========================================================

@app.get("/analysis/labor/{labor_code}")
def analyze_labor(
    labor_code: str,
    project_code: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    site_condition: str | None = None,
    session: Session = Depends(get_session),
):

    # -----------------------------------------------------
    # DATE VALIDATION
    # -----------------------------------------------------

    if start_date:
        validate_date_string(
            start_date,
            "start_date",
        )

    if end_date:
        validate_date_string(
            end_date,
            "end_date",
        )

    if (
        start_date
        and end_date
        and start_date > end_date
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "start_date must not be "
                "later than end_date."
            ),
        )

    # -----------------------------------------------------
    # LABOR MASTER
    # -----------------------------------------------------

    labor_item = session.exec(
        select(LaborMaster).where(
            LaborMaster.code
            == labor_code
        )
    ).first()

    if not labor_item:
        raise HTTPException(
            status_code=404,
            detail=(
                "Labor master code not found."
            ),
        )

    # -----------------------------------------------------
    # COMPANY SETTING
    # -----------------------------------------------------

    company_setting = session.exec(
        select(CompanySetting).where(
            CompanySetting.setting_code
            == "MINUTES_PER_PERSON_DAY"
        )
    ).first()

    if not company_setting:
        raise HTTPException(
            status_code=404,
            detail=(
                "MINUTES_PER_PERSON_DAY "
                "company setting not found."
            ),
        )

    if company_setting.value <= 0:
        raise HTTPException(
            status_code=500,
            detail=(
                "MINUTES_PER_PERSON_DAY "
                "must be greater than zero."
            ),
        )

    # -----------------------------------------------------
    # FILTER RECORDS
    # -----------------------------------------------------

    statement = select(
        ActualWorkRecord
    ).where(
        ActualWorkRecord.labor_code
        == labor_code
    )

    if project_code:
        statement = statement.where(
            ActualWorkRecord.project_code
            == project_code
        )

    if start_date:
        statement = statement.where(
            ActualWorkRecord.work_date
            >= start_date
        )

    if end_date:
        statement = statement.where(
            ActualWorkRecord.work_date
            <= end_date
        )

    if site_condition:
        statement = statement.where(
            ActualWorkRecord.site_condition
            == site_condition
        )

    records = session.exec(
        statement
    ).all()

    if not records:
        raise HTTPException(
            status_code=404,
            detail=(
                "No actual work records found."
            ),
        )

    # -----------------------------------------------------
    # CALCULATE VALID RECORDS
    # -----------------------------------------------------

    calculated_records = []

    for record in records:

        if record.quantity <= 0:
            continue

        if record.workers <= 0:
            continue

        if record.work_minutes <= 0:
            continue

        person_days = (
            record.workers
            * record.work_minutes
            / company_setting.value
        )

        labor_per_unit = (
            person_days
            / record.quantity
        )

        calculated_records.append(
            {
                "id": record.id,
                "project_code":
                    record.project_code,
                "work_date":
                    record.work_date,
                "site_condition":
                    record.site_condition,
                "quantity":
                    record.quantity,
                "actual_person_days":
                    person_days,
                "actual_labor_per_unit":
                    labor_per_unit,
            }
        )

    if not calculated_records:
        raise HTTPException(
            status_code=400,
            detail=(
                "No valid actual work records "
                "found."
            ),
        )

    values = [
        row[
            "actual_labor_per_unit"
        ]
        for row in calculated_records
    ]

    # -----------------------------------------------------
    # OUTLIER DETECTION
    # -----------------------------------------------------

    q1 = None
    q3 = None
    iqr = None
    lower_bound = None
    upper_bound = None

    outliers = []
    clean_records = calculated_records.copy()

    if len(values) >= 4:

        quartiles = quantiles(
            values,
            n=4,
            method="inclusive",
        )

        q1 = quartiles[0]
        q3 = quartiles[2]

        iqr = q3 - q1

        lower_bound = (
            q1
            - OUTLIER_IQR_MULTIPLIER
            * iqr
        )

        upper_bound = (
            q3
            + OUTLIER_IQR_MULTIPLIER
            * iqr
        )

        outliers = [
            row
            for row in calculated_records
            if (
                row[
                    "actual_labor_per_unit"
                ] < lower_bound
                or row[
                    "actual_labor_per_unit"
                ] > upper_bound
            )
        ]

        outlier_ids = {
            row["id"]
            for row in outliers
        }

        clean_records = [
            row
            for row in calculated_records
            if row["id"]
            not in outlier_ids
        ]

    # -----------------------------------------------------
    # CLEAN ANALYSIS VALUES
    # -----------------------------------------------------

    clean_values = [
        row[
            "actual_labor_per_unit"
        ]
        for row in clean_records
    ]

    if not clean_values:
        raise HTTPException(
            status_code=400,
            detail=(
                "All records were classified "
                "as outliers."
            ),
        )

    total_quantity = sum(
        row["quantity"]
        for row in clean_records
    )

    total_person_days = sum(
        row["actual_person_days"]
        for row in clean_records
    )

    average_actual_labor = mean(
        clean_values
    )

    median_actual_labor = median(
        clean_values
    )

    minimum_actual_labor = min(
        clean_values
    )

    maximum_actual_labor = max(
        clean_values
    )

    weighted_actual_labor = (
        total_person_days
        / total_quantity
    )

    # -----------------------------------------------------
    # STANDARD LABOR
    # -----------------------------------------------------

    standard_labor_per_unit = (
        labor_item.standard_minutes
        / company_setting.value
    )

    difference = (
        weighted_actual_labor
        - standard_labor_per_unit
    )

    difference_percent = (
        difference
        / standard_labor_per_unit
        * 100
    )

    # -----------------------------------------------------
    # SAMPLE QUALITY WARNING
    # -----------------------------------------------------

    sample_count = len(
        clean_records
    )

    if (
        sample_count
        < MIN_ANALYSIS_SAMPLE_COUNT
    ):
        sample_status = "INSUFFICIENT"

        sample_warning = (
            "Sample count is too small "
            "for reliable standard labor "
            "decision."
        )

    elif (
        sample_count
        < RECOMMENDED_ANALYSIS_SAMPLE_COUNT
    ):
        sample_status = "CAUTION"

        sample_warning = (
            "Analysis is possible, but "
            "additional samples are "
            "recommended."
        )

    else:
        sample_status = "OK"

        sample_warning = None

    # -----------------------------------------------------
    # RESPONSE
    # -----------------------------------------------------

    return {
        "labor_code":
            labor_item.code,

        "labor_name":
            labor_item.name,

        "unit":
            labor_item.unit,

        "filters": {
            "project_code":
                project_code,
            "start_date":
                start_date,
            "end_date":
                end_date,
            "site_condition":
                site_condition,
        },

        "data_quality": {
            "raw_record_count":
                len(calculated_records),

            "analysis_record_count":
                len(clean_records),

            "outlier_count":
                len(outliers),

            "sample_status":
                sample_status,

            "sample_warning":
                sample_warning,
        },

        "outlier_detection": {
            "method":
                "IQR",

            "multiplier":
                OUTLIER_IQR_MULTIPLIER,

            "q1":
                round(q1, 4)
                if q1 is not None
                else None,

            "q3":
                round(q3, 4)
                if q3 is not None
                else None,

            "iqr":
                round(iqr, 4)
                if iqr is not None
                else None,

            "lower_bound":
                round(lower_bound, 4)
                if lower_bound
                is not None
                else None,

            "upper_bound":
                round(upper_bound, 4)
                if upper_bound
                is not None
                else None,

            "outliers":
                outliers,
        },

        "analysis": {
            "total_quantity":
                round(
                    total_quantity,
                    4,
                ),

            "total_actual_person_days":
                round(
                    total_person_days,
                    4,
                ),

            "average_actual_labor_per_unit":
                round(
                    average_actual_labor,
                    4,
                ),

            "weighted_actual_labor_per_unit":
                round(
                    weighted_actual_labor,
                    4,
                ),

            "median_actual_labor_per_unit":
                round(
                    median_actual_labor,
                    4,
                ),

            "minimum_actual_labor_per_unit":
                round(
                    minimum_actual_labor,
                    4,
                ),

            "maximum_actual_labor_per_unit":
                round(
                    maximum_actual_labor,
                    4,
                ),
        },

        "standard": {
            "standard_minutes":
                labor_item.standard_minutes,

            "minutes_per_person_day":
                company_setting.value,

            "standard_labor_per_unit":
                round(
                    standard_labor_per_unit,
                    4,
                ),

            "difference_from_standard":
                round(
                    difference,
                    4,
                ),

            "difference_percent":
                round(
                    difference_percent,
                    2,
                ),
        },
    }


# =========================================================
# STANDARD REVISION EVALUATION
# =========================================================

def evaluate_standard_revision(
    labor_code: str,
    session: Session,
):

    analysis_result = analyze_labor(
        labor_code=labor_code,
        project_code=None,
        start_date=None,
        end_date=None,
        site_condition=None,
        session=session,
    )

    data_quality = analysis_result[
        "data_quality"
    ]

    analysis = analysis_result[
        "analysis"
    ]

    standard = analysis_result[
        "standard"
    ]

    sample_count = data_quality[
        "analysis_record_count"
    ]

    raw_record_count = data_quality[
        "raw_record_count"
    ]

    outlier_count = data_quality[
        "outlier_count"
    ]

    weighted_actual = analysis[
        "weighted_actual_labor_per_unit"
    ]

    median_actual = analysis[
        "median_actual_labor_per_unit"
    ]

    difference_percent = standard[
        "difference_percent"
    ]

    minutes_per_person_day = standard[
        "minutes_per_person_day"
    ]

    # -----------------------------------------------------
    # CONSISTENCY CHECK
    # -----------------------------------------------------

    if median_actual > 0:
        mean_median_gap_percent = (
            abs(
                weighted_actual
                - median_actual
            )
            / median_actual
            * 100
        )
    else:
        mean_median_gap_percent = 999.0

    # -----------------------------------------------------
    # OUTLIER RATIO
    # -----------------------------------------------------

    if raw_record_count > 0:
        outlier_ratio = (
            outlier_count
            / raw_record_count
        )
    else:
        outlier_ratio = 1.0

    # -----------------------------------------------------
    # ELIGIBILITY CHECK
    # -----------------------------------------------------

    checks = {
        "sample_count_ok":
            sample_count
            >= REVISION_MIN_SAMPLE_COUNT,

        "difference_ok":
            abs(difference_percent)
            >= REVISION_DIFFERENCE_THRESHOLD_PERCENT,

        "consistency_ok":
            mean_median_gap_percent
            <= REVISION_MAX_MEAN_MEDIAN_GAP_PERCENT,

        "outlier_ratio_ok":
            outlier_ratio
            <= REVISION_MAX_OUTLIER_RATIO,
    }

    eligible = all(
        checks.values()
    )

    reasons = []

    if not checks["sample_count_ok"]:
        reasons.append(
            "Insufficient valid sample count."
        )

    if not checks["difference_ok"]:
        reasons.append(
            "Difference from current standard "
            "is below revision threshold."
        )

    if not checks["consistency_ok"]:
        reasons.append(
            "Weighted average and median "
            "are not sufficiently consistent."
        )

    if not checks["outlier_ratio_ok"]:
        reasons.append(
            "Outlier ratio is too high."
        )

    # -----------------------------------------------------
    # PROPOSED STANDARD
    # -----------------------------------------------------

    proposed_labor_per_unit = (
        weighted_actual
    )

    proposed_standard_minutes = (
        proposed_labor_per_unit
        * minutes_per_person_day
    )

    return {
        "labor_code":
            labor_code,

        "eligible":
            eligible,

        "checks":
            checks,

        "quality": {
            "sample_count":
                sample_count,

            "raw_record_count":
                raw_record_count,

            "outlier_count":
                outlier_count,

            "outlier_ratio":
                round(
                    outlier_ratio,
                    4,
                ),

            "mean_median_gap_percent":
                round(
                    mean_median_gap_percent,
                    2,
                ),
        },

        "current_standard": {
            "standard_minutes":
                standard[
                    "standard_minutes"
                ],

            "standard_labor_per_unit":
                standard[
                    "standard_labor_per_unit"
                ],
        },

        "actual_analysis": {
            "weighted_actual_labor_per_unit":
                weighted_actual,

            "median_actual_labor_per_unit":
                median_actual,

            "difference_percent":
                difference_percent,
        },

        "proposal": {
            "proposed_labor_per_unit":
                round(
                    proposed_labor_per_unit,
                    4,
                ),

            "proposed_standard_minutes":
                round(
                    proposed_standard_minutes,
                    2,
                ),
        },

        "reasons":
            reasons,
    }

# =========================================================
# STANDARD REVISION PREVIEW
# =========================================================

@app.get(
    "/standard-revision-preview/{labor_code}"
)
def preview_standard_revision(
    labor_code: str,
    session: Session = Depends(get_session),
):
    return evaluate_standard_revision(
        labor_code,
        session,
    )

# =========================================================
# CREATE STANDARD REVISION CANDIDATE
# =========================================================

@app.post(
    "/standard-revision-candidates/{labor_code}"
)
def create_standard_revision_candidate(
    labor_code: str,
    session: Session = Depends(get_session),
):

    labor_item = session.exec(
        select(LaborMaster).where(
            LaborMaster.code == labor_code
        )
    ).first()

    if not labor_item:
        raise HTTPException(
            status_code=404,
            detail="Labor master code not found.",
        )

    evaluation = evaluate_standard_revision(
        labor_code,
        session,
    )

    if not evaluation["eligible"]:
        raise HTTPException(
            status_code=400,
            detail={
                "message":
                    "Revision candidate is not eligible.",

                "evaluation":
                    evaluation,
            },
        )

    # -----------------------------------------------------
    # PENDING DUPLICATE CHECK
    # -----------------------------------------------------

    existing_pending = session.exec(
        select(StandardRevisionCandidate).where(
            StandardRevisionCandidate.labor_code
            == labor_code,

            StandardRevisionCandidate.status
            == "PENDING",
        )
    ).first()

    if existing_pending:
        raise HTTPException(
            status_code=409,
            detail=(
                "Pending revision candidate "
                f"already exists. ID: "
                f"{existing_pending.id}"
            ),
        )

    quality = evaluation[
        "quality"
    ]

    actual_analysis = evaluation[
        "actual_analysis"
    ]

    proposal = evaluation[
        "proposal"
    ]

    company_setting = session.exec(
        select(CompanySetting).where(
            CompanySetting.setting_code
            == "MINUTES_PER_PERSON_DAY"
        )
    ).first()

    candidate = StandardRevisionCandidate(
        labor_code=labor_code,

        current_revision=
            labor_item.revision,

        current_standard_minutes=
            labor_item.standard_minutes,

        proposed_standard_minutes=
            proposal[
                "proposed_standard_minutes"
            ],

        minutes_per_person_day=
            company_setting.value,

        proposed_labor_per_unit=
            proposal[
                "proposed_labor_per_unit"
            ],

        sample_count=
            quality["sample_count"],

        outlier_count=
            quality["outlier_count"],

        weighted_actual_labor_per_unit=
            actual_analysis[
                "weighted_actual_labor_per_unit"
            ],

        median_actual_labor_per_unit=
            actual_analysis[
                "median_actual_labor_per_unit"
            ],

        difference_percent=
            actual_analysis[
                "difference_percent"
            ],

        status="PENDING",

        reason=(
            "Generated from validated "
            "actual work analysis."
        ),
    )

    session.add(candidate)
    session.commit()
    session.refresh(candidate)

    return candidate

# =========================================================
# READ STANDARD REVISION CANDIDATES
# =========================================================

@app.get("/standard-revision-candidates")
def read_standard_revision_candidates(
    labor_code: str | None = None,
    status: str | None = None,
    session: Session = Depends(get_session),
):

    statement = select(
        StandardRevisionCandidate
    )

    if labor_code:
        statement = statement.where(
            StandardRevisionCandidate.labor_code
            == labor_code
        )

    if status:
        statement = statement.where(
            StandardRevisionCandidate.status
            == status
        )

    statement = statement.order_by(
        StandardRevisionCandidate.id
    )

    items = session.exec(
        statement
    ).all()

    return {
        "count": len(items),
        "records": items,
    }

# =========================================================
# APPROVE STANDARD REVISION
# =========================================================

@app.post(
    "/standard-revision-candidates/"
    "{candidate_id}/approve"
)
def approve_standard_revision(
    candidate_id: int,
    review: StandardRevisionReviewRequest,
    session: Session = Depends(get_session),
):

    candidate = session.get(
        StandardRevisionCandidate,
        candidate_id,
    )

    if not candidate:
        raise HTTPException(
            status_code=404,
            detail=(
                "Revision candidate not found."
            ),
        )

    if candidate.status != "PENDING":
        raise HTTPException(
            status_code=409,
            detail=(
                "Revision candidate is not pending."
            ),
        )

    labor_item = session.exec(
        select(LaborMaster).where(
            LaborMaster.code
            == candidate.labor_code
        )
    ).first()

    if not labor_item:
        raise HTTPException(
            status_code=404,
            detail="Labor master code not found.",
        )

    # -----------------------------------------------------
    # STALE CANDIDATE PROTECTION
    # -----------------------------------------------------

    if (
        labor_item.revision
        != candidate.current_revision
    ):
        raise HTTPException(
            status_code=409,
            detail=(
                "Labor master revision changed "
                "after candidate creation."
            ),
        )

    old_standard_minutes = (
        labor_item.standard_minutes
    )

    old_revision = labor_item.revision

    # -----------------------------------------------------
    # UPDATE OFFICIAL LABOR MASTER
    # -----------------------------------------------------

    labor_item.standard_minutes = (
        candidate.proposed_standard_minutes
    )

    labor_item.revision += 1

    labor_item.updated_at = datetime.now()

        # -----------------------------------------------------
    # UPDATE CANDIDATE
    # -----------------------------------------------------

    candidate.status = "APPROVED"
    candidate.reviewer = review.reviewer
    candidate.review_comment = review.comment
    candidate.reviewed_at = datetime.now()

    # -----------------------------------------------------
    # CREATE AUDIT LOG
    # -----------------------------------------------------

    audit_log = AuditLog(
        action="STANDARD_REVISION_APPROVED",
        entity_type="LaborMaster",
        entity_id=labor_item.id,
        labor_code=labor_item.code,
        candidate_id=candidate.id,
        actor=review.reviewer,
        old_standard_minutes=old_standard_minutes,
        new_standard_minutes=candidate.proposed_standard_minutes,
        old_revision=old_revision,
        new_revision=labor_item.revision,
        old_status="PENDING",
        new_status="APPROVED",
        comment=review.comment,
    )

    # -----------------------------------------------------
    # SAVE
    # -----------------------------------------------------

    session.add(labor_item)
    session.add(candidate)
    session.add(audit_log)

    session.commit()

    session.refresh(labor_item)
    session.refresh(candidate)
    session.refresh(labor_item)
    session.refresh(candidate)

    return {
        "message":
            "Standard revision approved.",

        "labor_code":
            labor_item.code,

        "old_revision":
            old_revision,

        "new_revision":
            labor_item.revision,

        "old_standard_minutes":
            old_standard_minutes,

        "new_standard_minutes":
            labor_item.standard_minutes,

        "candidate_id":
            candidate.id,

        "reviewer":
            candidate.reviewer,
    }


# =========================================================
# REJECT STANDARD REVISION
# =========================================================

@app.post(
    "/standard-revision-candidates/"
    "{candidate_id}/reject"
)
def reject_standard_revision(
    candidate_id: int,
    review: StandardRevisionReviewRequest,
    session: Session = Depends(get_session),
):

    candidate = session.get(
        StandardRevisionCandidate,
        candidate_id,
    )

    if not candidate:
        raise HTTPException(
            status_code=404,
            detail="Revision candidate not found.",
        )

    if candidate.status != "PENDING":
        raise HTTPException(
            status_code=409,
            detail="Revision candidate is not pending.",
        )

    # -----------------------------------------------------
    # UPDATE CANDIDATE
    # -----------------------------------------------------

    candidate.status = "REJECTED"
    candidate.reviewer = review.reviewer
    candidate.review_comment = review.comment
    candidate.reviewed_at = datetime.now()

    # -----------------------------------------------------
    # CREATE AUDIT LOG
    # -----------------------------------------------------

    audit_log = AuditLog(
        action="STANDARD_REVISION_REJECTED",
        entity_type="StandardRevisionCandidate",
        entity_id=candidate.id,
        labor_code=candidate.labor_code,
        candidate_id=candidate.id,
        actor=review.reviewer,
        old_standard_minutes=(
            candidate.current_standard_minutes
        ),
        new_standard_minutes=(
            candidate.current_standard_minutes
        ),
        old_revision=(
            candidate.current_revision
        ),
        new_revision=(
            candidate.current_revision
        ),
        old_status="PENDING",
        new_status="REJECTED",
        comment=review.comment,
    )

    # -----------------------------------------------------
    # SAVE
    # -----------------------------------------------------

    session.add(candidate)
    session.add(audit_log)

    session.commit()
    session.refresh(candidate)

    # -----------------------------------------------------
    # RESPONSE
    # -----------------------------------------------------

    return {
        "message": "Standard revision rejected.",
        "candidate_id": candidate.id,
        "labor_code": candidate.labor_code,
        "status": candidate.status,
        "reviewer": candidate.reviewer,
    }

# =========================================================
# AUDIT LOG
# =========================================================

@app.get("/audit-logs")
def read_audit_logs(
    labor_code: str | None = None,
    action: str | None = None,
    actor: str | None = None,
    session: Session = Depends(get_session),
):

    statement = select(AuditLog)

    if labor_code:
        statement = statement.where(
            AuditLog.labor_code == labor_code
        )

    if action:
        statement = statement.where(
            AuditLog.action == action
        )

    if actor:
        statement = statement.where(
            AuditLog.actor == actor
        )

    statement = statement.order_by(
        AuditLog.created_at.desc(),
        AuditLog.id.desc(),
    )

    items = session.exec(statement).all()

    return {
        "count": len(items),
        "records": items,
    }


# =========================================================
# LABOR ESTIMATION ENGINE
# =========================================================

@app.post("/estimate/labor")
def calculate_labor_estimate(
    request: LaborEstimateRequest,
    session: Session = Depends(get_session),
):

    # -----------------------------------------------------
    # 1. LABOR MASTER
    # -----------------------------------------------------

    labor_item = session.exec(
        select(LaborMaster).where(
            LaborMaster.code
            == request.labor_code
        )
    ).first()

    if not labor_item:
        raise HTTPException(
            status_code=404,
            detail="Labor master code not found.",
        )

    # -----------------------------------------------------
    # 2. COMPANY SETTING
    # -----------------------------------------------------

    company_setting = session.exec(
        select(CompanySetting).where(
            CompanySetting.setting_code
            == "MINUTES_PER_PERSON_DAY"
        )
    ).first()

    if not company_setting:
        raise HTTPException(
            status_code=404,
            detail=(
                "MINUTES_PER_PERSON_DAY "
                "company setting not found."
            ),
        )

    # -----------------------------------------------------
    # 3. LABOR COST MASTER
    # -----------------------------------------------------

    labor_cost_item = session.exec(
        select(LaborCostMaster).where(
            LaborCostMaster.cost_code
            == request.cost_code
        )
    ).first()

    if not labor_cost_item:
        raise HTTPException(
            status_code=404,
            detail="Labor cost code not found.",
        )

    # -----------------------------------------------------
    # 4. VALIDATION
    # -----------------------------------------------------

    if request.quantity <= 0:
        raise HTTPException(
            status_code=400,
            detail="Quantity must be greater than zero.",
        )

    if request.correction_factor <= 0:
        raise HTTPException(
            status_code=400,
            detail=(
                "Correction factor must be "
                "greater than zero."
            ),
        )

    if company_setting.value <= 0:
        raise HTTPException(
            status_code=500,
            detail=(
                "MINUTES_PER_PERSON_DAY "
                "must be greater than zero."
            ),
        )

    if labor_cost_item.unit_price < 0:
        raise HTTPException(
            status_code=500,
            detail=(
                "Labor unit price "
                "cannot be negative."
            ),
        )

    # -----------------------------------------------------
    # 5. CALCULATION
    # -----------------------------------------------------

    standard_labor_per_unit = (
        labor_item.standard_minutes
        / company_setting.value
    )

    required_labor = (
        request.quantity
        * standard_labor_per_unit
        * request.correction_factor
    )

    total_labor_cost = (
        required_labor
        * labor_cost_item.unit_price
    )

    # -----------------------------------------------------
    # 6. RESPONSE
    # -----------------------------------------------------

    return {
        "labor_code":
            labor_item.code,

        "labor_name":
            labor_item.name,

        "unit":
            labor_item.unit,

        "quantity":
            request.quantity,

        "standard_minutes_per_unit":
            labor_item.standard_minutes,

        "minutes_per_person_day":
            company_setting.value,

        "standard_labor_per_unit":
            round(
                standard_labor_per_unit,
                4,
            ),

        "correction_factor":
            request.correction_factor,

        "required_labor":
            round(
                required_labor,
                4,
            ),

        "cost_code":
            labor_cost_item.cost_code,

        "role_name":
            labor_cost_item.role_name,

        "labor_unit_price":
            labor_cost_item.unit_price,

        "currency":
            labor_cost_item.currency,

        "labor_cost":
            round(
                total_labor_cost,
                2,
            ),
    }

@app.post("/material-master")
def create_material_master(
    item: MaterialMaster,
    session: Session = Depends(get_session),
):
    existing = session.exec(
        select(MaterialMaster).where(
            MaterialMaster.material_code == item.material_code
        )
    ).first()

    if existing:
        raise HTTPException(
            status_code=409,
            detail="This material code already exists.",
        )

    item.id = None
    item.created_at = datetime.now()
    item.updated_at = datetime.now()

    session.add(item)
    session.commit()
    session.refresh(item)

    return item

@app.get("/material-master")
def read_material_master(
    category: str | None = None,
    manufacturer: str | None = None,
    active: bool | None = None,
    session: Session = Depends(get_session),
):
    statement = select(MaterialMaster)

    if category:
        statement = statement.where(
            MaterialMaster.category == category
        )

    if manufacturer:
        statement = statement.where(
            MaterialMaster.manufacturer == manufacturer
        )

    if active is not None:
        statement = statement.where(
            MaterialMaster.active == active
        )

    statement = statement.order_by(
        MaterialMaster.material_code
    )

    return session.exec(statement).all()

@app.post("/estimate/combined")
def estimate_combined(
    request: CombinedEstimateRequest,
    session: Session = Depends(get_session),
):
    # 1. Material Master
    material = session.exec(
        select(MaterialMaster).where(
            MaterialMaster.material_code == request.material_code
        )
    ).first()

    if not material:
        raise HTTPException(
            status_code=404,
            detail="Material code not found.",
        )

    if not material.active:
        raise HTTPException(
            status_code=400,
            detail="Material is inactive.",
        )

    # 2. Labor Master
    labor = session.exec(
        select(LaborMaster).where(
            LaborMaster.code == request.labor_code
        )
    ).first()

    if not labor:
        raise HTTPException(
            status_code=404,
            detail="Labor code not found.",
        )

    if not labor.active:
        raise HTTPException(
            status_code=400,
            detail="Labor master is inactive.",
        )

    # 3. Labor Cost Master
    labor_cost = session.exec(
        select(LaborCostMaster).where(
            LaborCostMaster.cost_code == request.cost_code
        )
    ).first()

    if not labor_cost:
        raise HTTPException(
            status_code=404,
            detail="Labor cost code not found.",
        )

    if not labor_cost.active:
        raise HTTPException(
            status_code=400,
            detail="Labor cost master is inactive.",
        )

    # 4. Company Setting
    setting = session.exec(
        select(CompanySetting).where(
            CompanySetting.setting_code
            == "MINUTES_PER_PERSON_DAY"
        )
    ).first()

    if not setting:
        raise HTTPException(
            status_code=404,
            detail="MINUTES_PER_PERSON_DAY setting not found.",
        )

    minutes_per_person_day = setting.value

    # 5. Validation
    if request.quantity <= 0:
        raise HTTPException(
            status_code=400,
            detail="Quantity must be greater than zero.",
        )

    if request.correction_factor <= 0:
        raise HTTPException(
            status_code=400,
            detail="Correction factor must be greater than zero.",
        )

    if minutes_per_person_day <= 0:
        raise HTTPException(
            status_code=400,
            detail="MINUTES_PER_PERSON_DAY must be greater than zero.",
        )

    # 6. Material calculation
    material_cost = (
        material.unit_price
        * request.quantity
    )

    # 7. Labor calculation
    standard_labor_per_unit = (
        labor.standard_minutes
        / minutes_per_person_day
    )

    labor_person_days = (
        standard_labor_per_unit
        * request.quantity
        * request.correction_factor
    )

    labor_cost_total = (
        labor_person_days
        * labor_cost.unit_price
    )

    # 8. Direct cost
    direct_cost = (
        material_cost
        + labor_cost_total
    )

    return {
        "material": {
            "material_code": material.material_code,
            "material_name": material.material_name,
            "unit": material.unit,
            "unit_price": round(material.unit_price, 2),
            "quantity": request.quantity,
            "material_cost": round(material_cost, 2),
        },
        "labor": {
            "labor_code": labor.code,
            "labor_name": labor.name,
            "standard_minutes": round(
                labor.standard_minutes,
                2,
            ),
            "minutes_per_person_day": round(
                minutes_per_person_day,
                2,
            ),
            "standard_labor_per_unit": round(
                standard_labor_per_unit,
                4,
            ),
            "correction_factor": request.correction_factor,
            "labor_person_days": round(
                labor_person_days,
                4,
            ),
            "cost_code": labor_cost.cost_code,
            "labor_unit_price": round(
                labor_cost.unit_price,
                2,
            ),
            "labor_cost": round(
                labor_cost_total,
                2,
            ),
        },
        "summary": {
            "material_cost": round(
                material_cost,
                2,
            ),
            "labor_cost": round(
                labor_cost_total,
                2,
            ),
            "direct_cost": round(
                direct_cost,
                2,
            ),
            "currency": material.currency,
        },
    }

@app.post("/estimates")
def create_estimate(
    request: EstimateCreateRequest,
    session: Session = Depends(get_session),
):
    existing = session.exec(
        select(Estimate).where(
            Estimate.estimate_number
            == request.estimate_number
        )
    ).first()

    if existing:
        raise HTTPException(
            status_code=409,
            detail="This estimate number already exists.",
        )

    if request.overhead_rate < 0:
        raise HTTPException(
            status_code=400,
            detail="Overhead rate cannot be negative.",
        )

    if request.profit_rate < 0:
        raise HTTPException(
            status_code=400,
            detail="Profit rate cannot be negative.",
        )

    estimate = Estimate(
        estimate_number=request.estimate_number,
        project_code=request.project_code,
        customer_name=request.customer_name,
        estimate_title=request.estimate_title,
        status="DRAFT",
        currency=request.currency,
        material_cost_total=0,
        labor_cost_total=0,
        direct_cost_total=0,
        overhead_rate=request.overhead_rate,
        overhead_amount=0,
        profit_rate=request.profit_rate,
        profit_amount=0,
        estimate_total=0,
        description=request.description,
    )

    session.add(estimate)
    session.commit()
    session.refresh(estimate)

    return estimate

@app.get("/estimates")
def read_estimates(
    status: str | None = None,
    project_code: str | None = None,
    session: Session = Depends(get_session),
):
    statement = select(Estimate)

    if status:
        statement = statement.where(
            Estimate.status == status
        )

    if project_code:
        statement = statement.where(
            Estimate.project_code == project_code
        )

    statement = statement.order_by(
        Estimate.id.desc()
    )

    return session.exec(statement).all()

@app.post("/estimates/{estimate_id}/items")
def create_estimate_item(
    estimate_id: int,
    request: EstimateItemCreateRequest,
    session: Session = Depends(get_session),
):
    estimate = session.get(Estimate, estimate_id)

    if not estimate:
        raise HTTPException(
            status_code=404,
            detail="Estimate not found.",
        )

    if estimate.status != "DRAFT":
        raise HTTPException(
            status_code=409,
            detail="Only DRAFT estimates can be edited.",
        )

    if request.quantity <= 0:
        raise HTTPException(
            status_code=400,
            detail="Quantity must be greater than zero.",
        )

    if request.correction_factor <= 0:
        raise HTTPException(
            status_code=400,
            detail="Correction factor must be greater than zero.",
        )

    material = session.exec(
        select(MaterialMaster).where(
            MaterialMaster.material_code == request.material_code
        )
    ).first()

    if not material:
        raise HTTPException(
            status_code=404,
            detail="Material code not found.",
        )

    labor = session.exec(
        select(LaborMaster).where(
            LaborMaster.code == request.labor_code
        )
    ).first()

    if not labor:
        raise HTTPException(
            status_code=404,
            detail="Labor code not found.",
        )

    labor_cost = session.exec(
        select(LaborCostMaster).where(
            LaborCostMaster.cost_code == request.cost_code
        )
    ).first()

    if not labor_cost:
        raise HTTPException(
            status_code=404,
            detail="Labor cost code not found.",
        )

    setting = session.exec(
        select(CompanySetting).where(
            CompanySetting.setting_code
            == "MINUTES_PER_PERSON_DAY"
        )
    ).first()

    if not setting:
        raise HTTPException(
            status_code=404,
            detail="MINUTES_PER_PERSON_DAY setting not found.",
        )

    existing_items = session.exec(
        select(EstimateItem).where(
            EstimateItem.estimate_id == estimate_id
        )
    ).all()

    line_number = len(existing_items) + 1

    material_cost = (
        material.unit_price
        * request.quantity
    )

    standard_labor_per_unit = (
        labor.standard_minutes
        / setting.value
    )

    labor_person_days = (
        standard_labor_per_unit
        * request.quantity
        * request.correction_factor
    )

    labor_cost_total = (
        labor_person_days
        * labor_cost.unit_price
    )

    direct_cost = (
        material_cost
        + labor_cost_total
    )

    item = EstimateItem(
        estimate_id=estimate_id,
        line_number=line_number,

        material_code=material.material_code,
        labor_code=labor.code,
        cost_code=labor_cost.cost_code,

        description=request.description,

        quantity=request.quantity,
        unit=material.unit,

        material_unit_price=material.unit_price,
        material_cost=material_cost,

        standard_minutes=labor.standard_minutes,
        correction_factor=request.correction_factor,
        labor_person_days=labor_person_days,

        labor_unit_price=labor_cost.unit_price,
        labor_cost=labor_cost_total,

        direct_cost=direct_cost,
    )

    session.add(item)
    session.commit()
    session.refresh(item)

    items = session.exec(
        select(EstimateItem).where(
            EstimateItem.estimate_id == estimate_id
        )
    ).all()

    material_cost_total = sum(
        x.material_cost for x in items
    )

    labor_cost_total_sum = sum(
        x.labor_cost for x in items
    )

    direct_cost_total = (
        material_cost_total
        + labor_cost_total_sum
    )

    overhead_amount = (
        direct_cost_total
        * estimate.overhead_rate
        / 100
    )

    subtotal = (
        direct_cost_total
        + overhead_amount
    )

    profit_amount = (
        subtotal
        * estimate.profit_rate
        / 100
    )

    estimate.material_cost_total = material_cost_total
    estimate.labor_cost_total = labor_cost_total_sum
    estimate.direct_cost_total = direct_cost_total
    estimate.overhead_amount = overhead_amount
    estimate.profit_amount = profit_amount
    estimate.estimate_total = subtotal + profit_amount
    estimate.updated_at = datetime.now()

    session.add(estimate)
    session.commit()

    session.refresh(item)
    session.refresh(estimate)

    return {
    "item": item,
    "estimate": estimate,
}

@app.delete("/estimates/{estimate_id}/items/{item_id}")
def delete_estimate_item(
    estimate_id: int,
    item_id: int,
    session: Session = Depends(get_session),
):
    estimate = session.get(Estimate, estimate_id)

    if not estimate:
        raise HTTPException(
            status_code=404,
            detail="Estimate not found.",
        )

    if estimate.status != "DRAFT":
        raise HTTPException(
            status_code=409,
            detail="Only DRAFT estimates can be edited.",
        )

    item = session.get(EstimateItem, item_id)

    if not item:
        raise HTTPException(
            status_code=404,
            detail="Estimate item not found.",
        )

    if item.estimate_id != estimate_id:
        raise HTTPException(
            status_code=409,
            detail="Estimate item does not belong to this estimate.",
        )

    session.delete(item)
    session.commit()

    items = session.exec(
        select(EstimateItem).where(
            EstimateItem.estimate_id == estimate_id
        )
    ).all()

    material_cost_total = sum(
        x.material_cost for x in items
    )

    labor_cost_total = sum(
        x.labor_cost for x in items
    )

    direct_cost_total = (
        material_cost_total
        + labor_cost_total
    )

    overhead_amount = (
        direct_cost_total
        * estimate.overhead_rate
        / 100
    )

    subtotal = (
        direct_cost_total
        + overhead_amount
    )

    profit_amount = (
        subtotal
        * estimate.profit_rate
        / 100
    )

    estimate.material_cost_total = material_cost_total
    estimate.labor_cost_total = labor_cost_total
    estimate.direct_cost_total = direct_cost_total
    estimate.overhead_amount = overhead_amount
    estimate.profit_amount = profit_amount
    estimate.estimate_total = subtotal + profit_amount
    estimate.updated_at = datetime.now()

    session.add(estimate)
    session.commit()
    session.refresh(estimate)

    return {
        "message": "Estimate item deleted.",
        "deleted_item_id": item_id,
        "estimate": estimate,
    }

@app.get("/estimates/{estimate_id}/items")
def read_estimate_items(
    estimate_id: int,
    session: Session = Depends(get_session),
):
    estimate = session.get(Estimate, estimate_id)

    if not estimate:
        raise HTTPException(
            status_code=404,
            detail="Estimate not found.",
        )

    items = session.exec(
        select(EstimateItem)
        .where(
            EstimateItem.estimate_id == estimate_id
        )
        .order_by(
            EstimateItem.line_number
        )
    ).all()

    return {
        "estimate_id": estimate_id,
        "estimate_number": estimate.estimate_number,
        "count": len(items),
        "items": items,
    }

@app.put("/estimates/{estimate_id}/items/{item_id}")
def update_estimate_item(
    estimate_id: int,
    item_id: int,
    request: EstimateItemUpdateRequest,
    session: Session = Depends(get_session),
):
    estimate = session.get(Estimate, estimate_id)

    if not estimate:
        raise HTTPException(
            status_code=404,
            detail="Estimate not found.",
        )

    if estimate.status != "DRAFT":
        raise HTTPException(
            status_code=409,
            detail="Only DRAFT estimates can be edited.",
        )

    item = session.get(EstimateItem, item_id)

    if not item:
        raise HTTPException(
            status_code=404,
            detail="Estimate item not found.",
        )

    if item.estimate_id != estimate_id:
        raise HTTPException(
            status_code=409,
            detail="Estimate item does not belong to this estimate.",
        )

    if request.quantity is not None:
        if request.quantity <= 0:
            raise HTTPException(
                status_code=400,
                detail="Quantity must be greater than zero.",
            )
        item.quantity = request.quantity

    if request.correction_factor is not None:
        if request.correction_factor <= 0:
            raise HTTPException(
                status_code=400,
                detail="Correction factor must be greater than zero.",
            )
        item.correction_factor = request.correction_factor

    if request.description is not None:
        item.description = request.description

    setting = session.exec(
        select(CompanySetting).where(
            CompanySetting.setting_code
            == "MINUTES_PER_PERSON_DAY"
        )
    ).first()

    if not setting:
        raise HTTPException(
            status_code=404,
            detail="MINUTES_PER_PERSON_DAY setting not found.",
        )

    if setting.value <= 0:
        raise HTTPException(
            status_code=400,
            detail="MINUTES_PER_PERSON_DAY must be greater than zero.",
        )

    item.material_cost = (
        item.material_unit_price
        * item.quantity
    )

    standard_labor_per_unit = (
        item.standard_minutes
        / setting.value
    )

    item.labor_person_days = (
        standard_labor_per_unit
        * item.quantity
        * item.correction_factor
    )

    item.labor_cost = (
        item.labor_person_days
        * item.labor_unit_price
    )

    item.direct_cost = (
        item.material_cost
        + item.labor_cost
    )

    session.add(item)
    session.commit()
    session.refresh(item)

    items = session.exec(
        select(EstimateItem).where(
            EstimateItem.estimate_id == estimate_id
        )
    ).all()

    material_cost_total = sum(
        x.material_cost for x in items
    )

    labor_cost_total = sum(
        x.labor_cost for x in items
    )

    direct_cost_total = (
        material_cost_total
        + labor_cost_total
    )

    overhead_amount = (
        direct_cost_total
        * estimate.overhead_rate
        / 100
    )

    subtotal = (
        direct_cost_total
        + overhead_amount
    )

    profit_amount = (
        subtotal
        * estimate.profit_rate
        / 100
    )

    estimate.material_cost_total = material_cost_total
    estimate.labor_cost_total = labor_cost_total
    estimate.direct_cost_total = direct_cost_total
    estimate.overhead_amount = overhead_amount
    estimate.profit_amount = profit_amount
    estimate.estimate_total = subtotal + profit_amount
    estimate.updated_at = datetime.now()

    session.add(estimate)
    session.commit()

    session.refresh(item)
    session.refresh(estimate)

    return {
        "item": item,
        "estimate": estimate,
    }

@app.put("/estimates/{estimate_id}/status")
def update_estimate_status(
    estimate_id: int,
    request: EstimateStatusUpdateRequest,
    session: Session = Depends(get_session),
):
    estimate = session.get(Estimate, estimate_id)

    if not estimate:
        raise HTTPException(
            status_code=404,
            detail="Estimate not found.",
        )

    allowed_statuses = {
        "DRAFT",
        "APPROVED",
        "ISSUED",
    }

    new_status = request.status.upper()

    if new_status not in allowed_statuses:
        raise HTTPException(
            status_code=400,
            detail="Invalid estimate status.",
        )

    current_status = estimate.status

    allowed_transitions = {
        "DRAFT": {"APPROVED"},
        "APPROVED": {"ISSUED"},
        "ISSUED": set(),
    }

    if new_status == current_status:
        return {
            "estimate": estimate,
            "message": "Estimate status is unchanged.",
        }

    if new_status not in allowed_transitions.get(
        current_status,
        set(),
    ):
        raise HTTPException(
            status_code=409,
            detail=(
                f"Status transition from "
                f"{current_status} to {new_status} "
                f"is not allowed."
            ),
        )

    estimate.status = new_status
    estimate.updated_at = datetime.now()

    session.add(estimate)
    session.commit()
    session.refresh(estimate)

    return {
        "estimate": estimate,
        "message": (
            f"Estimate status changed "
            f"from {current_status} "
            f"to {new_status}."
        ),
    }

@app.get("/estimates/{estimate_id}/export/excel")
def export_estimate_excel(
    estimate_id: int,
    session: Session = Depends(get_session),
):
    estimate = session.get(Estimate, estimate_id)

    if not estimate:
        raise HTTPException(
            status_code=404,
            detail="Estimate not found.",
        )

    items = session.exec(
        select(EstimateItem)
        .where(
            EstimateItem.estimate_id == estimate_id
        )
        .order_by(
            EstimateItem.line_number
        )
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Estimate"

    # -----------------------------------------------------
    # PAGE SETTINGS
    # -----------------------------------------------------

    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 16

    # -----------------------------------------------------
    # STYLES
    # -----------------------------------------------------

    title_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_fill = PatternFill(
        "solid",
        fgColor="D9EAF7",
    )

    total_fill = PatternFill(
        "solid",
        fgColor="E2F0D9",
    )

    thin = Side(
        style="thin",
        color="808080",
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    ws.merge_cells("A1:H2")

    ws["A1"] = "ESTIMATE"

    ws["A1"].font = Font(
        size=20,
        bold=True,
        color="FFFFFF",
    )

    ws["A1"].fill = title_fill

    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    # -----------------------------------------------------
    # ESTIMATE HEADER
    # -----------------------------------------------------

    ws["A4"] = "Estimate No."
    ws["B4"] = estimate.estimate_number

    ws["A5"] = "Customer"
    ws["B5"] = estimate.customer_name

    ws["A6"] = "Project"
    ws["B6"] = estimate.estimate_title

    ws["A7"] = "Project Code"
    ws["B7"] = estimate.project_code

    ws["E4"] = "Status"
    ws["F4"] = estimate.status

    ws["E5"] = "Currency"
    ws["F5"] = estimate.currency

    ws["E6"] = "Created"
    ws["F6"] = estimate.created_at

    for cell in [
        "A4",
        "A5",
        "A6",
        "A7",
        "E4",
        "E5",
        "E6",
    ]:
        ws[cell].font = Font(
            bold=True
        )

    # -----------------------------------------------------
    # DETAIL HEADER
    # -----------------------------------------------------

    header_row = 10

    headers = [
        "No.",
        "Description",
        "Qty",
        "Unit",
        "Material",
        "Labor",
        "Direct Cost",
        "Labor Days",
    ]

    for col, value in enumerate(
        headers,
        start=1,
    ):
        cell = ws.cell(
            row=header_row,
            column=col,
            value=value,
        )

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill
        cell.border = border

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # -----------------------------------------------------
    # DETAIL ROWS
    # -----------------------------------------------------

    current_row = header_row + 1

    for item in items:

        values = [
            item.line_number,
            item.description,
            item.quantity,
            item.unit,
            item.material_cost,
            item.labor_cost,
            item.direct_cost,
            item.labor_person_days,
        ]

        for col, value in enumerate(
            values,
            start=1,
        ):
            cell = ws.cell(
                row=current_row,
                column=col,
                value=value,
            )

            cell.border = border

        current_row += 1

    # -----------------------------------------------------
    # NUMBER FORMAT
    # -----------------------------------------------------

    for row in range(
        header_row + 1,
        current_row,
    ):
        ws.cell(
            row=row,
            column=3,
        ).number_format = '#,##0.00'

        for col in [
            5,
            6,
            7,
        ]:
            ws.cell(
                row=row,
                column=col,
            ).number_format = '#,##0'

        ws.cell(
            row=row,
            column=8,
        ).number_format = '0.000'

    # -----------------------------------------------------
    # TOTALS
    # -----------------------------------------------------

    total_row = current_row + 1

    totals = [
        (
            "Material Total",
            estimate.material_cost_total,
        ),
        (
            "Labor Total",
            estimate.labor_cost_total,
        ),
        (
            "Direct Cost",
            estimate.direct_cost_total,
        ),
        (
            f"Overhead ({estimate.overhead_rate}%)",
            estimate.overhead_amount,
        ),
        (
            f"Profit ({estimate.profit_rate}%)",
            estimate.profit_amount,
        ),
        (
            "ESTIMATE TOTAL",
            estimate.estimate_total,
        ),
    ]

    for label, value in totals:

        ws.merge_cells(
            start_row=total_row,
            start_column=5,
            end_row=total_row,
            end_column=6,
        )

        label_cell = ws.cell(
            row=total_row,
            column=5,
            value=label,
        )

        value_cell = ws.cell(
            row=total_row,
            column=7,
            value=value,
        )

        label_cell.font = Font(
            bold=True
        )

        value_cell.font = Font(
            bold=True
        )

        value_cell.number_format = '#,##0'

        for col in range(
            5,
            8,
        ):
            ws.cell(
                row=total_row,
                column=col,
            ).border = border

        if label == "ESTIMATE TOTAL":

            for col in range(
                5,
                8,
            ):
                ws.cell(
                    row=total_row,
                    column=col,
                ).fill = total_fill

            value_cell.font = Font(
                bold=True,
                size=14,
            )

        total_row += 1

    # -----------------------------------------------------
    # NOTES
    # -----------------------------------------------------

    if estimate.description:

        ws["A" + str(total_row + 2)] = (
            "Notes"
        )

        ws["A" + str(total_row + 2)].font = Font(
            bold=True
        )

        ws.merge_cells(
            start_row=total_row + 3,
            start_column=1,
            end_row=total_row + 5,
            end_column=8,
        )

        note_cell = ws.cell(
            row=total_row + 3,
            column=1,
            value=estimate.description,
        )

        note_cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    # -----------------------------------------------------
    # OUTPUT
    # -----------------------------------------------------

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    filename = (
        f"{estimate.estimate_number}.xlsx"
    )

    return StreamingResponse(
        output,
        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
            f'attachment; filename="{filename}"'
        },
    )