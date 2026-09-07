import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlmodel import Session, SQLModel, create_engine

from backend.database import engine, get_session
from backend.main import app
from backend.models import EstimateDocument, QuoteItem, WorkTypeMaster


TEST_ENGINE = engine
SQLModel.metadata.create_all(TEST_ENGINE)


def _override_get_session():
    with Session(TEST_ENGINE) as session:
        yield session


app.dependency_overrides[get_session] = _override_get_session
client = TestClient(app)


def _reset_db():
    SQLModel.metadata.drop_all(TEST_ENGINE)
    SQLModel.metadata.create_all(TEST_ENGINE)


def _seed_document_and_item():
    with Session(TEST_ENGINE) as session:
        document = EstimateDocument(
            source_filename="sample.pdf",
            file_hash="hash-analysis-export",
            source_type="VENDOR_QUOTE",
            time_category="CURRENT",
            vendor_name="Example Vendor",
            project_code="PRJ-001",
            quote_number="Q-001",
            verification_status="VERIFIED",
        )
        session.add(document)
        session.commit()
        session.refresh(document)

        work_type = WorkTypeMaster(
            work_type_code="MASONRY",
            category="STRUCTURE",
            sub_category="WALL",
            work_type_name="Masonry Work",
            default_unit="EA",
        )
        session.add(work_type)
        session.commit()
        session.refresh(work_type)

        item = QuoteItem(
            document_id=document.id,
            line_number=1,
            raw_description="Masonry wall panel installation",
            normalized_description="Masonry wall panel installation",
            work_type_code=work_type.work_type_code,
            quantity=2.0,
            unit="EA",
            amount=30000.0,
            source_unit_price=15000.0,
            calculated_unit_price=15000.0,
            normalized_unit_price=15000.0,
            mapping_confidence=0.9,
            verification_status="VERIFIED",
        )
        session.add(item)
        session.commit()
        session.refresh(item)

        return document.id, item.id


def test_document_analysis_export_returns_xlsx_download():
    _reset_db()
    document_id, _ = _seed_document_and_item()

    response = client.get(f"/documents/{document_id}/analysis/export/excel")

    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]
    assert "attachment; filename=" in response.headers["content-disposition"]
    assert len(response.content) > 0

    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == [
        "SUMMARY",
        "SOURCE_DOCUMENT",
        "QUOTE_ITEMS",
        "WORK_TYPE_ANALYSIS",
        "PRICE_STATISTICS",
        "LABOR_ANALYSIS",
        "OUTLIERS",
        "RECOMMENDATIONS",
        "VERIFICATION",
        "AUDIT_INFO",
    ]
    assert workbook["SUMMARY"]["A1"].value == "CConstructionAI"


def test_document_analysis_export_404_for_missing_document():
    _reset_db()
    response = client.get("/documents/999999/analysis/export/excel")

    assert response.status_code == 404
    assert response.json()["detail"] == "Estimate document not found."
